# -*- coding: utf-8 -*-
"""1주차 원본 PDF 의 모든 문장이 결과물에 남아 있는지 기계적으로 확인한다.

글자·숫자만 남기고(글머리표·구분기호·문장부호는 뺀다) 부분문자열로 대조한다.
누락이 하나라도 있으면 종료 코드 1 을 낸다.

    python3 verify-coverage.py <원본 1주차.pdf>
"""
import re, sys, unicodedata
import pymupdf
from pptx import Presentation

# 원본 1주차 PDF 경로 — 첫 인자로 넘기거나 아래 기본값을 고친다
PDF  = sys.argv[1] if len(sys.argv) > 1 else "원본_1주차.pdf"
import os
HERE = os.path.dirname(os.path.abspath(__file__))
TGTS = {"PPTX": f"{HERE}/week01.pptx", "HTML": f"{HERE}/week01.html", "MD": f"{HERE}/week01.md"}

MAP = str.maketrans({
    "·":"·","‧":"·","〮":"·","•":"·","∙":"·","・":"·","．":".","，":",",
    "‘":"'","’":"'","“":'"','”':'"',"–":"-","—":"-","―":"-","−":"-","~":"~","～":"~",
    "（":"(","）":")","：":":","；":";","／":"/","％":"%","！":"!","？":"?",
    "¼":"4",
})
WORD = re.compile(r"[0-9A-Za-z\uac00-\ud7a3\u4e00-\u9fff@]+")
def norm(t):
    """글자·숫자만 남긴다 — 글머리표(➢ ✓ -), 구분기호(_ / :), 문장부호 차이는
       '내용이 같은가' 와 무관하므로 비교에서 뺀다."""
    t = unicodedata.normalize("NFKC", t).translate(MAP)
    return "".join(WORD.findall(t)).lower()

LEADNUM = re.compile(r"^\d{1,2}")   # '1. 오리엔테이션' 처럼 표 앞머리 번호가 다른 칸에 있는 경우

# ── 원본 PDF 문장 ────────────────────────────────────────────────────
doc = pymupdf.open(PDF)
lines, seen = [], set()
for pno, page in enumerate(doc, 1):
    for raw in page.get_text().splitlines():
        n = norm(raw)
        if len(n) < 2:            continue
        if n in ("chapter","c","h","a","p","t","e","r"):  continue
        if n.isdigit():           continue          # 쪽번호
        if n in seen:             continue
        seen.add(n)
        lines.append((pno, raw.strip(), n))

# ── 결과물 텍스트 ────────────────────────────────────────────────────
def pptx_text(p):
    prs, out = Presentation(p), []
    for sl in prs.slides:
        for sp in sl.shapes:
            if sp.has_text_frame: out.append(sp.text_frame.text)
            if sp.has_table:
                for r in sp.table.rows:
                    for c in r.cells: out.append(c.text)
            if sp.has_chart:
                ch = sp.chart
                for pl in ch.plots:
                    out.extend(list(pl.categories))
        if sl.has_notes_slide: out.append(sl.notes_slide.notes_text_frame.text)
    return "\n".join(out)

def file_text(p):
    s = open(p, encoding="utf-8").read()
    s = re.sub(r"<script\b.*?</script>", " ", s, flags=re.S|re.I)
    s = re.sub(r"<style\b.*?</style>", " ", s, flags=re.S|re.I)
    s = re.sub(r"<!--.*?-->", " ", s, flags=re.S)
    s = re.sub(r"<[^>]+>", " ", s)
    for a, b in [("&lsquo;","'"),("&rsquo;","'"),("&nbsp;"," "),("&amp;","&"),
                 ("&frac14;","¼"),("&#8226;","·"),("&lt;","<"),("&gt;",">")]:
        s = s.replace(a, b)
    s = re.sub(r"&#\d+;|&[a-z]+;", " ", s)
    return s

blobs = {k: norm(pptx_text(v) if v.endswith(".pptx") else file_text(v)) for k, v in TGTS.items()}

print(f"원본 PDF 문장 {len(lines)}개 대조\n" + "=" * 62)
bad = {}
for name, blob in blobs.items():
    def has(n):
        if n in blob: return True
        n2 = LEADNUM.sub("", n)          # 앞 번호를 뗀 형태도 인정
        return len(n2) > 3 and n2 in blob
    miss = [(p, raw) for p, raw, n in lines if not has(n)]
    hit  = len(lines) - len(miss)
    print(f"{name:5s}  {hit}/{len(lines)}  = {hit/len(lines)*100:6.2f} %")
    if miss: bad[name] = miss
for name, miss in bad.items():
    print(f"\n── {name} 누락 {len(miss)}건 ──")
    for p, raw in miss[:40]:
        print(f"  p{p:02d} | {raw[:88]}")
sys.exit(1 if bad else 0)
